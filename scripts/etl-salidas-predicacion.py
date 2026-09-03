#!/usr/bin/env python3
"""Importa el XLSX de salidas únicamente al staging de Supabase.

Uso (ensayo, no necesita credenciales):

    python3 scripts/etl-salidas-predicacion.py \
      --input /tmp/Salidas-de-predicacion.xlsx --dry-run

Uso real contra el clon (la service role key se lee de un archivo):

    SUPABASE_PROJECT_REF=rkmioktcsgqqjshrlkmy \
    SUPABASE_SERVICE_ROLE_KEY_FILE=/ruta/a/clave-del-clon.txt \
    python3 scripts/etl-salidas-predicacion.py \
      --input /tmp/Salidas-de-predicacion.xlsx

El script crea una corrida en ``importaciones`` y registros en
``importacion_registros``. Nunca inserta en salidas, salida_resultados,
territorio_historial ni conductores. ``--dry-run`` valida la clasificación y
los números esperados sin escribir nada.

La planilla debe ser la copia leída desde Drive con ID
1L_VEocJhVt-Qykqs4sWkZO7QIWVwjjHc. La identidad de cada corrida se controla
con SHA-256 del archivo y versión del parser. Se conserva la fila completa:
cada celda incluye su valor de origen, fórmula, valor calculado (si lo hay),
formato, comentario, combinación y enlace. Un error de fórmula invalida la
normalización de esa fila, pero nunca se descarta el bruto.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen

warnings.filterwarnings(
    "ignore",
    message="DrawingML support is incomplete and limited to charts and images only.*",
    category=UserWarning,
)

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - mensaje para otra máquina
    raise SystemExit(
        "Falta openpyxl. Usá el runtime de planillas de Workspace o instalalo "
        "en el entorno de Python que ejecute este script."
    ) from exc


DEV_REF = "rkmioktcsgqqjshrlkmy"
PROD_REF = "dwgvzcnarrjgqjotocdw"
DRIVE_ID = "1L_VEocJhVt-Qykqs4sWkZO7QIWVwjjHc"
SOURCE_FILE = "Salidas de predicación.xlsx"
PARSER_VERSION = "2.0.0-temporal"

HISTORY_COLOR_MEANING = {
    "FFFFFF00": "manana",
    "FFFF9900": "tarde",
    "FFCCCCCC": "territorio_personal",
    "FF980000": "superintendente_circuito",
    "FF00FFFF": "grupo_1",
    "FF00FF00": "grupo_2",
    "FF4A86E8": "grupo_3",
    "FF8E7CC3": "grupo_4",
    "FFD5A6BD": "grupo_5",
}

EXPECTED_COUNTS = {
    "salida": 1854,
    "historial_territorio": 1031,
    "conductor": 55,
    "grupo": 5,
    "punto_encuentro": 191,
    "territorio": 200,
    "territorio_personal": 9,
    "otro": 1,
}

EXPECTED_AGENDA_BY_SHEET = {
    "Ultima salida vigente": 48,
    "REGISTRO 2007 AL 1508": 64,
    "REGISTRO 2206 AL 1807": 52,
    "REGISTRO 2505 AL 2006": 52,
    "REGISTRO 2704 AL 2305": 52,
    "REGISTRO 3003 AL 2504": 52,
    "REGISTRO 0203 AL 2903 + SALIDA ": 50,
    "REGISTRO 0202 AL 2802": 52,
    "REGISTO 0501 AL 3101": 52,
    "REGISTRO 081225 A 030126": 46,
    "REGISTRO 2411 A 0712": 26,
    "REGISTRO 1011 A 2311": 26,
    "REGISTRO 2710 A 911": 26,
    "REGISTRO 1310 A 2610": 26,
    "REGISTRO 2909 A 1210": 25,
    "Salida SC 09-25": 12,
    "REGISTRO 0809 A 2109": 26,
    "REGISTRO 2508 A 0709": 25,
    "REGISTRO 1108 A 2408": 27,
    "REGISTRO 2807 A 1008": 25,
    "REGISTRO 1407 a 2707": 27,
    "REGISTRO 3006 A 1307": 26,
    "REGISTRO 1606 A 2906": 26,
    "REGISTRO 0206 A 1506": 25,
    "REGISTRO 1905 A 0106": 26,
    "REGISTRO 0505 A 1805": 26,
    "REGISTRO 2104 A 0405": 26,
    "REGISTRO 704 a 2004": 26,
    "REGISTRO 24-03 al 06-04": 26,
    "Salida SC": 12,
    "REGISTRO 03 a 1603": 26,
    "REGISTRO 1702 a 0203": 26,
    "REGISTRO 0302 a 1602": 26,
    "REGISTRO 2001 a 0202": 26,
    "REGISTRO601 a 1901": 26,
    "REGISTRO 2312 a 0501": 19,
    "REGISTRO 0912 A 2112": 25,
    "REGISTRO 2511 a 0712": 25,
    "REGISTRO 1111 A 2311": 25,
    "REGISTRO 2810 a 0911": 26,
    "REGISTRO  14 A 26 de octubre": 22,
    "Registro territorios unificado": 265,
    "REGISTRO 30 de septiembre al 12": 26,
    "REGISTRO 12 al 24 de septiembre": 26,
    "REGISTRO 27 de agosto al 10 de ": 25,
    "REGISTRO 15 al 27 de agosto": 24,
    "REGISTRO 31J a 13 Agosto": 23,
    "REGISTRO 18 al 30 de julio": 24,
    "REGISTRO 4 al 16 de julio": 24,
    "REGISTRO 2O de junio a 2 de jul": 24,
    "REGISTRO 6 al 18 de junio ": 24,
    "REGISTRO 23 de mayo a 4 de juni": 24,
    "REGISTRO 9 a 21 de mayo": 24,
    "REGISTRO 24 de abril al 8 de ma": 24,
    "REGISTRO 15 a 22 de abril": 15,
}

HORIZONTAL_EXTRA_BLOCKS = (
    ("REGISTRO 2312 a 0501", 19, 9, 16),
    ("REGISTRO 2312 a 0501", 19, 17, 24),
)

# `importacion_registros` tiene una restricción única por pestaña, fila y
# tipo, pero el historial del Excel tiene varias entradas (B:Q) en la misma
# fila. La fila visible se conserva en bruto; este identificador estable solo
# evita colisiones en staging sin tocar la migración.
HISTORY_SYNTHETIC_ROW_FACTOR = 1000


class ETLError(RuntimeError):
    """Error de lectura, clasificación o carga del ETL."""


def normalize_text(value: Any) -> str:
    """Normaliza texto solo para comparar encabezados, nunca para guardar alias."""

    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ").strip()
    text = unicodedata.normalize("NFKD", text)
    return "".join(char for char in text if not unicodedata.combining(char)).lower()


def json_value(value: Any) -> Any:
    """Convierte tipos de Excel a valores JSON sin perder la forma visible."""

    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return str(value)
        return value
    if isinstance(value, str):
        return value
    return str(value)


def is_formula(cell: Any) -> bool:
    return cell.data_type == "f" or (
        isinstance(cell.value, str) and cell.value.startswith("=")
    )


def looks_like_excel_error(value: Any) -> bool:
    return isinstance(value, str) and value.lstrip().startswith("#")


def is_formula_error(source_cell: Any, calculated_cell: Any) -> bool:
    if source_cell.data_type == "e" or calculated_cell.data_type == "e":
        return True
    return looks_like_excel_error(source_cell.value) or looks_like_excel_error(
        calculated_cell.value
    )


def cell_kind(source_cell: Any, calculated_cell: Any, formula_error: bool) -> str:
    if formula_error:
        return "error_formula"
    if is_formula(source_cell):
        return "formula"
    value = source_cell.value
    if value is None:
        return "vacio"
    if isinstance(value, bool):
        return "booleano"
    if isinstance(value, datetime) or isinstance(value, date):
        return "fecha"
    if isinstance(value, time):
        return "hora"
    if isinstance(value, (int, float)):
        return "numero"
    if looks_like_excel_error(value):
        return "error"
    return "texto"


class WorkbookReader:
    def __init__(self, path: Path):
        if not path.is_file():
            raise ETLError(f"No existe el XLSX de entrada: {path}")
        self.path = path
        self.source_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
        self.source_size_bytes = path.stat().st_size
        self.source = load_workbook(path, data_only=False, keep_links=True, read_only=False)
        self.calculated = load_workbook(
            path, data_only=True, keep_links=True, read_only=False
        )
        self._raw_cache: dict[tuple[str, int], dict[str, Any]] = {}
        self._merged_by_cell: dict[tuple[str, int, int], str] = {}
        self._merged_anchor: dict[tuple[str, int, int], tuple[int, int]] = {}
        for worksheet in self.source.worksheets:
            for merged in worksheet.merged_cells.ranges:
                merged_ref = str(merged)
                anchor = (merged.min_row, merged.min_col)
                for merged_row in range(merged.min_row, merged.max_row + 1):
                    for merged_column in range(merged.min_col, merged.max_col + 1):
                        key = (worksheet.title, merged_row, merged_column)
                        self._merged_by_cell[key] = merged_ref
                        self._merged_anchor[key] = anchor
        self.used_columns = {
            ws.title: self._used_column_count(ws) for ws in self.source.worksheets
        }

    @staticmethod
    def _used_column_count(ws: Any) -> int:
        maximum = 1
        for row in ws.iter_rows():
            for cell in row:
                if (
                    cell.value is not None
                    or cell.hyperlink is not None
                    or cell.comment is not None
                ):
                    maximum = max(maximum, cell.column)
        return maximum

    def source_cell(self, sheet: str, row: int, column: int) -> Any:
        return self.source[sheet].cell(row, column)

    def calculated_cell(self, sheet: str, row: int, column: int) -> Any:
        return self.calculated[sheet].cell(row, column)

    def effective(self, sheet: str, row: int, column: int) -> Any:
        source_cell = self.source_cell(sheet, row, column)
        calculated_cell = self.calculated_cell(sheet, row, column)
        if is_formula(source_cell):
            if is_formula_error(source_cell, calculated_cell):
                return None
            return calculated_cell.value
        return source_cell.value

    def effective_in_merge(self, sheet: str, row: int, column: int | None) -> Any:
        """Lee la celda o, sólo si pertenece a un merge, su ancla real."""

        if column is None:
            return None
        anchor = self._merged_anchor.get((sheet, row, column))
        if anchor is None:
            return self.effective(sheet, row, column)
        return self.effective(sheet, anchor[0], anchor[1])

    @staticmethod
    def _fill_payload(cell: Any) -> dict[str, Any] | None:
        fill = cell.fill
        if fill is None or fill.fill_type is None:
            return None

        def color_payload(color: Any) -> dict[str, Any] | None:
            if color is None or color.type is None:
                return None
            payload = {"tipo": color.type}
            value = getattr(color, color.type, None)
            if value is not None:
                payload["valor"] = value
            if color.tint:
                payload["tint"] = color.tint
            return payload

        return {
            "tipo": fill.fill_type,
            "frente": color_payload(fill.fgColor),
            "fondo": color_payload(fill.bgColor),
        }

    def raw_row(self, sheet: str, row: int) -> dict[str, Any]:
        key = (sheet, row)
        if key in self._raw_cache:
            return self._raw_cache[key]

        width = self.used_columns[sheet]
        cells: list[dict[str, Any]] = []
        formula_errors: list[str] = []

        for column in range(1, width + 1):
            source_cell = self.source_cell(sheet, row, column)
            calculated_cell = self.calculated_cell(sheet, row, column)
            formula = is_formula(source_cell)
            formula_error = is_formula_error(source_cell, calculated_cell)
            letter = get_column_letter(column)
            item: dict[str, Any] = {
                "columna": letter,
                "valor": json_value(source_cell.value),
                "tipo": cell_kind(source_cell, calculated_cell, formula_error),
                "formato": source_cell.number_format,
                "texto_origen": None if source_cell.value is None else str(source_cell.value),
            }
            if formula:
                item["formula"] = json_value(source_cell.value)
                item["valor_calculado"] = json_value(calculated_cell.value)
            if source_cell.hyperlink is not None:
                item["hipervinculo"] = source_cell.hyperlink.target
            if source_cell.comment is not None:
                item["comentario"] = {
                    "texto": source_cell.comment.text,
                    "autor": source_cell.comment.author,
                }
            fill = self._fill_payload(source_cell)
            if fill is not None:
                item["relleno"] = fill
            merged_ref = self._merged_by_cell.get((sheet, row, column))
            if merged_ref is not None:
                item["rango_combinado"] = merged_ref
            if formula_error:
                formula_errors.append(letter)
            cells.append(item)

        raw: dict[str, Any] = {
            "hoja": sheet,
            "fila": row,
            "archivo_sha256": self.source_sha256,
            "parser_version": PARSER_VERSION,
            "celdas": cells,
        }
        if formula_errors:
            raw["errores_formula"] = formula_errors
        self._raw_cache[key] = raw
        return raw

    def raw_with_context(
        self,
        sheet: str,
        row: int,
        *,
        cell: str | None = None,
        subregistro: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        raw = copy(self.raw_row(sheet, row))
        if cell is not None:
            raw["celda_origen"] = cell
        if subregistro is not None:
            raw.update(subregistro)
        return raw

    def row_has_formula_error(self, sheet: str, row: int) -> bool:
        return bool(self.formula_error_columns(sheet, row))

    def formula_error_columns(self, sheet: str, row: int) -> list[str]:
        errors: list[str] = []
        for column in range(1, self.used_columns[sheet] + 1):
            source_cell = self.source_cell(sheet, row, column)
            calculated_cell = self.calculated_cell(sheet, row, column)
            if is_formula_error(source_cell, calculated_cell):
                errors.append(get_column_letter(column))
        return errors

    def close(self) -> None:
        self.source.close()
        self.calculated.close()


def is_time_value(value: Any) -> bool:
    if isinstance(value, time):
        return True
    return isinstance(value, datetime) and value.time() != time(0, 0)


def normalize_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        candidate = value.date()
    elif isinstance(value, date):
        candidate = value
    elif isinstance(value, str):
        text = value.strip()
        try:
            candidate = date.fromisoformat(text[:10])
        except ValueError:
            return None
    else:
        return None

    # 1900-01-01 en una fórmula es el cache de una fecha no resuelta, no un
    # dato histórico que debamos inventar.
    if candidate.year < 1901:
        return None
    return candidate.isoformat()


def normalize_time(value: Any) -> str | None:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, str):
        text = value.strip()
        for pattern in ("%H:%M:%S", "%H:%M"):
            try:
                return datetime.strptime(text, pattern).time().isoformat()
            except ValueError:
                pass
    return None


def header_name(value: Any) -> str:
    return normalize_text(value).replace(".", "")


def is_agenda_sheet(title: str) -> bool:
    if title in {"Modelo", "REGISTRO TERRITORIOS"}:
        return False
    normalized = normalize_text(title)
    return (
        title == "Ultima salida vigente"
        or normalized.startswith("registro")
        or normalized.startswith("registo")
        or normalized.startswith("salida")
    )


def is_agenda_header(values: list[Any]) -> bool:
    headers = [header_name(value) for value in values]
    has_date = "fecha" in headers
    has_time = "hora" in headers
    has_territory = any(value == "terr" for value in headers)
    has_driver = "conductor" in headers
    return has_date and has_time and (has_territory or has_driver)


def make_header_info(values: list[Any], row_number: int) -> dict[str, Any]:
    headers = [header_name(value) for value in values]

    def first(predicate: Any) -> int | None:
        for index, value in enumerate(headers, start=1):
            if predicate(value):
                return index
        return None

    date_col = first(lambda value: value == "fecha")
    time_col = first(lambda value: value == "hora")
    territory_col = first(lambda value: value == "terr")
    driver_col = first(lambda value: value == "conductor")
    meeting_col = first(
        lambda value: value in {"punto de encuentro", "punto encuentro"}
    )
    place_col = first(lambda value: value == "lugar")
    priority_col = first(lambda value: value == "priorizar")
    status_col = first(lambda value: value.startswith("territorio completado"))
    day_col = first(lambda value: value == "dia")

    # En enero de 2026 el encabezado perdió la palabra Conductor. El
    # conductor sigue siendo la columna inmediatamente anterior al punto.
    if driver_col is None and meeting_col is not None and meeting_col > 1:
        driver_col = meeting_col - 1

    narrative_cols = []
    for index, value in enumerate(headers, start=1):
        if any(token in value for token in ("predicad", "observacion", "informacion")):
            narrative_cols.append(index)

    return {
        "header_row": row_number,
        "date_col": date_col,
        "time_col": time_col,
        "territory_col": territory_col,
        "driver_col": driver_col,
        "meeting_col": meeting_col,
        "place_col": place_col,
        "priority_col": priority_col,
        "status_col": status_col,
        "day_col": day_col,
        "narrative_cols": narrative_cols,
        "headers": values,
    }


def special_no_hour_row(reader: WorkbookReader, sheet: str, row: int) -> bool:
    # Las cuatro filas conocidas son salidas reales sin hora. El predicado
    # mantiene la regla semántica: B tiene una fecha y C no tiene hora; AC
    # o SIN PREDICACION identifica el registro y no una nota de la hoja.
    b = reader.effective(sheet, row, 2)
    c = reader.effective(sheet, row, 3)
    d = reader.effective(sheet, row, 4)
    f = reader.effective(sheet, row, 6)
    if b is None or is_time_value(c):
        return False
    d_text = normalize_text(d)
    return f is not None or "sin predicacion" in d_text


def agenda_events(reader: WorkbookReader) -> tuple[list[dict[str, Any]], dict[str, int]]:
    events: list[dict[str, Any]] = []
    counts: dict[str, int] = Counter()

    for sheet in reader.source.sheetnames:
        if not is_agenda_sheet(sheet):
            continue
        ws = reader.source[sheet]
        width = reader.used_columns[sheet]
        active_info: dict[str, Any] | None = None

        for row in range(1, ws.max_row + 1):
            values = [ws.cell(row, column).value for column in range(1, width + 1)]
            if is_agenda_header(values):
                active_info = make_header_info(values, row)
                continue
            if active_info is None:
                continue

            hour_value = None
            if active_info["time_col"] is not None:
                hour_value = reader.effective(sheet, row, active_info["time_col"])
            is_vertical = is_time_value(hour_value)
            is_special = special_no_hour_row(reader, sheet, row)
            if not is_vertical and not is_special:
                continue

            event = {"sheet": sheet, "row": row, "info": copy(active_info)}
            events.append(event)
            counts[sheet] += 1

    return events, dict(counts)


def classify_salida_type(value: Any) -> str | None:
    if value is None:
        return None
    text = normalize_text(value).replace(" ", "")
    if text == "tel" or "telefon" in text:
        return "telefonica"
    if text in {"ac", "ar"} or "asamblea" in text:
        return "asamblea"
    if text in {"zo", "zoom"}:
        return "especial"
    if "grupo" in text:
        return "grupos"
    return None


def source_and_effective(
    reader: WorkbookReader, sheet: str, row: int, column: int | None
) -> tuple[Any, Any]:
    if column is None:
        return None, None
    source = reader.source_cell(sheet, row, column).value
    effective = reader.effective(sheet, row, column)
    return source, effective


def source_code_text(value: Any) -> str | None:
    """Los códigos son identificadores léxicos, no cantidades calculables."""

    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def territory_references(value: Any) -> list[dict[str, Any]]:
    """Conserva todas las referencias explícitas sin resolverlas contra el mapa actual."""

    original = source_code_text(value)
    if original is None:
        return []
    parts = [part.strip() for part in re.split(r"\s*(?:,|/|\+|\by\b)\s*", original, flags=re.I)]
    parts = [part for part in parts if part]
    if len(parts) == 1:
        parts = [original]
    return [
        {
            "codigo_fuente": part,
            "rol": "principal" if index == 0 else "secundario",
            "orden": index,
            "estado_resolucion": "sin_resolver",
        }
        for index, part in enumerate(parts)
    ]


def fill_rgb(cell: Any) -> str | None:
    color = cell.fill.fgColor
    if cell.fill.fill_type != "solid" or color.type != "rgb":
        return None
    return color.rgb


def normalize_agenda_event(
    reader: WorkbookReader,
    event: dict[str, Any],
    *,
    horizontal: tuple[int, int] | None = None,
) -> tuple[dict[str, Any] | None, str | None, str | None]:
    sheet = event["sheet"]
    row = event["row"]
    info = event["info"]

    if horizontal is None:
        mapping = {
            "day_col": info["day_col"],
            "date_col": info["date_col"],
            "time_col": info["time_col"],
            "place_col": info["place_col"],
            "type_col": info["territory_col"],
            "territory_col": info["territory_col"],
            "driver_col": info["driver_col"],
            "meeting_col": info["meeting_col"],
            "priority_col": info["priority_col"],
            "status_col": info["status_col"],
            "narrative_cols": info["narrative_cols"],
        }
        range_start = 1
        range_end = reader.used_columns[sheet]
        extra = None
    else:
        range_start, range_end = horizontal
        mapping = {
            "day_col": range_start,
            "date_col": range_start + 1,
            "time_col": range_start + 2,
            "place_col": range_start + 3,
            "type_col": range_start + 4,
            "territory_col": range_start + 5,
            "driver_col": range_start + 6,
            "meeting_col": None,
            "priority_col": range_start + 4,
            "status_col": range_start + 7,
            "narrative_cols": [],
        }
        extra = {
            "fila_origen": row,
            "subregistro": {
                "rango": f"{get_column_letter(range_start)}{row}:{get_column_letter(range_end)}{row}"
            },
        }

    def effective(field: str) -> Any:
        column = mapping[field]
        return reader.effective(sheet, row, column) if column is not None else None

    # Día y fecha suelen estar combinados para varias salidas. Sólo se heredan
    # desde el ancla cuando la celda pertenece realmente al mismo merge.
    day = reader.effective_in_merge(sheet, row, mapping["day_col"])
    date_value = reader.effective_in_merge(sheet, row, mapping["date_col"])
    time_value = effective("time_col")
    place = effective("place_col")
    type_value = effective("type_col")
    territory = effective("territory_col")
    driver = effective("driver_col")
    meeting = effective("meeting_col")
    priority = effective("priority_col")
    status = effective("status_col")

    status_state: str
    if isinstance(status, bool):
        # Una casilla desmarcada no prueba que la salida no se haya realizado:
        # también se usó como pendiente/no informado. La negación necesita una
        # decisión humana o evidencia narrativa explícita.
        status_state = "realizada" if status else "sin_dato"
    else:
        # Un status ausente o vacío es información faltante, no una
        # negación. Esta rama cuenta las 106 salidas sin booleano concluyente.
        status_state = "sin_dato"

    normalized_date = normalize_date(date_value)
    normalized_time = normalize_time(time_value)
    reasons: list[str] = []

    if reader.row_has_formula_error(sheet, row):
        columns = ", ".join(reader.raw_row(sheet, row).get("errores_formula", []))
        reasons.append(f"error de fórmula en la fila (celdas: {columns})")

    if date_value is not None and normalized_date is None:
        reasons.append("fecha de origen no concluyente; se deja null")
    if time_value is not None and normalized_time is None:
        reasons.append("hora de origen no concluyente; se deja null")

    narratives: dict[str, Any] = {}
    for column in mapping["narrative_cols"]:
        label = reader.source_cell(sheet, info["header_row"], column).value
        narratives[f"{get_column_letter(column)}:{label}"] = json_value(
            reader.effective(sheet, row, column)
        )

    narrative_text = " ".join(
        str(value) for value in narratives.values() if value is not None
    )
    if status is True and re.search(
        r"\b(no|nunca|sin)\s+(se\s+)?predic|no\s+se\s+hizo|falt[oó]",
        normalize_text(narrative_text),
    ):
        reasons.append("el booleano indica realizada pero la narrativa indica que no se predicó")

    normalized: dict[str, Any] = {
        "fecha": normalized_date,
        "hora": normalized_time,
        "dia_bruto": json_value(day),
        "territorio_bruto": json_value(territory),
        "territorio_codigo_fuente": source_code_text(territory),
        "referencias_territoriales": territory_references(territory),
        "conductor_alias": json_value(driver),
        "lugar_bruto": json_value(place),
        "punto_encuentro_bruto": json_value(meeting),
        "priorizar_bruto": json_value(priority),
        "tipo_salida": classify_salida_type(type_value if horizontal else territory),
        "estado": status_state,
        "estado_fuente": json_value(status),
        "estado_resolucion": "sin_confirmar" if status is False else "resuelto",
        "narrativa": narratives,
    }
    if horizontal is not None:
        normalized["subregistro_rango"] = extra["subregistro"]["rango"]
        normalized["tipo_bloque_bruto"] = json_value(type_value)

    if reasons:
        return None, " ; ".join(reasons), "salidas"
    return normalized, None, "salidas"


def build_salida_records(
    reader: WorkbookReader,
    events: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for event in events:
        sheet = event["sheet"]
        row = event["row"]
        normalized, reason, destination = normalize_agenda_event(reader, event)
        records.append(
            {
                "pestania": sheet,
                "fila": row,
                "rango": f"A{row}:{get_column_letter(reader.used_columns[sheet])}{row}",
                "tipo": "salida",
                "bruto": reader.raw_row(sheet, row),
                "normalizado": normalized,
                "estado": "conflicto" if reason else "pendiente",
                "motivo": reason,
                "destino_tabla": destination,
            }
        )

    # La hoja contiene dos bloques horizontales en la fila 19. Son dos
    # registros del origen, pero la migración identifica por (pestaña, fila,
    # tipo). Se usa una fila sintética estable y se conserva fila_origen y el
    # rango exacto en bruto; no se toca la migración ni se fusionan datos.
    horizontal_sheet = "REGISTRO 2312 a 0501"
    vertical = next(
        (event for event in events if event["sheet"] == horizontal_sheet and event["row"] == 19),
        None,
    )
    if vertical is None:
        raise ETLError("No encontré la fila 19 de los bloques horizontales esperados.")

    for synthetic_row, (_, row, start, end) in zip((19001, 19002), HORIZONTAL_EXTRA_BLOCKS):
        normalized, reason, destination = normalize_agenda_event(
            reader, vertical, horizontal=(start, end)
        )
        bruto = reader.raw_with_context(
            horizontal_sheet,
            row,
            subregistro={
                "fila_origen": row,
                "subregistro": {
                    "rango": f"{get_column_letter(start)}{row}:{get_column_letter(end)}{row}"
                },
            },
        )
        records.append(
            {
                "pestania": horizontal_sheet,
                "fila": synthetic_row,
                "rango": f"{get_column_letter(start)}{row}:{get_column_letter(end)}{row}",
                "tipo": "salida",
                "bruto": bruto,
                "normalizado": normalized,
                "estado": "conflicto" if reason else "pendiente",
                "motivo": reason,
                "destino_tabla": destination,
            }
        )
    return records


def build_conductor_records(
    reader: WorkbookReader,
    events: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    seen: dict[str, tuple[dict[str, Any], int]] = {}
    for event in events:
        sheet = event["sheet"]
        row = event["row"]
        column = event["info"]["driver_col"]
        if column is None:
            continue
        source = reader.source_cell(sheet, row, column).value
        if source is None:
            continue
        alias = str(source)
        if alias not in seen:
            seen[alias] = (event, column)

    records: list[dict[str, Any]] = []
    for alias, (event, column) in seen.items():
        sheet = event["sheet"]
        row = event["row"]
        records.append(
            {
                "pestania": sheet,
                "fila": row,
                "rango": f"{get_column_letter(column)}{row}",
                "tipo": "conductor",
                "bruto": reader.raw_with_context(
                    sheet, row, cell=f"{get_column_letter(column)}{row}"
                ),
                "normalizado": {
                    "alias": alias,
                    "conductor_id": None,
                    "confianza": "dudoso",
                },
                "estado": "pendiente",
                "motivo": None,
                "destino_tabla": "conductores",
            }
        )
    return records


def build_history_records(reader: WorkbookReader) -> list[dict[str, Any]]:
    sheet = "REGISTRO TERRITORIOS"
    ws = reader.source[sheet]
    records: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        territory = reader.effective(sheet, row, 1)
        if territory is None:
            continue
        for column in range(2, 18):  # B:Q inclusive
            source = reader.source_cell(sheet, row, column).value
            value = reader.effective(sheet, row, column)
            if source is None and value is None:
                continue

            normalized: dict[str, Any] | None
            reason: str | None = None
            color_meaning = HISTORY_COLOR_MEANING.get(
                fill_rgb(reader.source_cell(sheet, row, column)) or ""
            )
            if reader.row_has_formula_error(sheet, row):
                normalized = None
                columns = ", ".join(reader.raw_row(sheet, row).get("errores_formula", []))
                reason = f"error de fórmula en la fila (celdas: {columns})"
            elif isinstance(value, (datetime, date)):
                normalized = {
                    "territorio_bruto": json_value(territory),
                    "territorio_codigo_fuente": source_code_text(territory),
                    "slot_columna": get_column_letter(column),
                    "fecha": normalize_date(value),
                    "texto": None,
                    "clase": "trabajado",
                    "valor_bruto": json_value(value),
                    "clasificacion_visual": color_meaning,
                }
                if normalized["fecha"] is None:
                    normalized = None
                    reason = "fecha de historial no concluyente; se deja null"
            elif isinstance(value, str) and normalize_text(value).startswith("bloq"):
                normalized = {
                    "territorio_bruto": json_value(territory),
                    "territorio_codigo_fuente": source_code_text(territory),
                    "slot_columna": get_column_letter(column),
                    "fecha": None,
                    "texto": value,
                    "clase": "bloqueo",
                    "valor_bruto": value,
                    "clasificacion_visual": color_meaning,
                }
            else:
                # Incluye el único texto de excepción (25-07). Los 70
                # textos permanecen texto; ninguno se parsea como fecha.
                normalized = None
                reason = "texto del historial no es una fecha; se conserva en bruto para revisión"

            records.append(
                {
                    "pestania": sheet,
                    "fila": row * HISTORY_SYNTHETIC_ROW_FACTOR + column,
                    "rango": f"{get_column_letter(column)}{row}",
                    "tipo": "historial_territorio",
                    "bruto": reader.raw_with_context(
                        sheet,
                        row,
                        cell=f"{get_column_letter(column)}{row}",
                        subregistro={"fila_origen": row},
                    ),
                    "normalizado": normalized,
                    "estado": "conflicto" if reason else "pendiente",
                    "motivo": reason,
                    "destino_tabla": "territorio_historial",
                }
            )
    return records


def is_google_maps_link(target: str) -> bool:
    lowered = target.lower()
    return any(
        token in lowered
        for token in (
            "google.com/maps",
            "maps.google.",
            "maps.app.goo.gl",
            "goo.gl/maps",
        )
    )


def build_territory_records(
    reader: WorkbookReader,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    sheet = "Territorios"
    ws = reader.source[sheet]
    candidates: list[int] = []
    for row in range(4, ws.max_row + 1):
        values = [ws.cell(row, column).value for column in range(1, 16)]
        if any(value is not None for value in values):
            candidates.append(row)

    code_counts = Counter(
        str(reader.effective(sheet, row, 2))
        for row in candidates
        if reader.effective(sheet, row, 2) is not None
    )
    territory_records: list[dict[str, Any]] = []
    point_records: list[dict[str, Any]] = []
    zoom_count = 0

    for row in candidates:
        code = reader.effective(sheet, row, 2)
        meeting = reader.effective(sheet, row, 3)
        neighborhood = reader.effective(sheet, row, 4)
        location_label = reader.effective(sheet, row, 5)
        source_e = reader.source_cell(sheet, row, 5)
        target = source_e.hyperlink.target if source_e.hyperlink is not None else None
        reasons: list[str] = []
        normalized: dict[str, Any] | None = {
            "codigo_bruto": json_value(code),
            "punto_encuentro_bruto": json_value(meeting),
            "barrio_bruto": json_value(neighborhood),
            "ubicacion_etiqueta": json_value(location_label),
        }

        if code_counts[str(code)] > 1:
            reasons.append(
                "código de territorio duplicado en Territorios; requiere revisión"
            )
            normalized = None
        if reader.row_has_formula_error(sheet, row):
            columns = ", ".join(reader.raw_row(sheet, row).get("errores_formula", []))
            reasons.append(f"error de fórmula en la fila (celdas: {columns})")
            normalized = None
        if target is not None and not is_google_maps_link(target):
            zoom_count += 1
            normalized = normalized or {
                "codigo_bruto": json_value(code),
                "punto_encuentro_bruto": json_value(meeting),
                "barrio_bruto": json_value(neighborhood),
                "ubicacion_etiqueta": json_value(location_label),
            }
            normalized["ubicacion_proveedor"] = "zoom"
            reasons.append(
                "enlace externo Zoom conservado solo en bruto; no se carga como punto de encuentro"
            )
        elif target is not None and normalized is not None:
            normalized["ubicacion_proveedor"] = "google_maps"

        territory_records.append(
            {
                "pestania": sheet,
                "fila": row,
                "rango": f"A{row}:{get_column_letter(reader.used_columns[sheet])}{row}",
                "tipo": "territorio",
                "bruto": reader.raw_row(sheet, row),
                "normalizado": normalized,
                "estado": "conflicto" if reasons else "pendiente",
                "motivo": " ; ".join(reasons) if reasons else None,
                "destino_tabla": "territorios",
            }
        )

        if target is not None and is_google_maps_link(target):
            point_records.append(
                {
                    "pestania": sheet,
                    "fila": row,
                    "rango": f"E{row}",
                    "tipo": "punto_encuentro",
                    "bruto": reader.raw_with_context(sheet, row, cell=f"E{row}"),
                    "normalizado": {
                        "territorio_bruto": json_value(code),
                        "etiqueta": json_value(location_label),
                        "url": target,
                        "proveedor": "google_maps",
                    },
                    "estado": "pendiente",
                    "motivo": None,
                    "destino_tabla": "puntos_encuentro",
                }
            )

    return territory_records, point_records, zoom_count


def build_group_records(reader: WorkbookReader) -> list[dict[str, Any]]:
    sheet = "Grupos - Colores"
    records: list[dict[str, Any]] = []
    for row in range(2, 7):
        name = reader.effective(sheet, row, 1)
        if name is None:
            continue
        territories = [
            json_value(reader.effective(sheet, row, column))
            for column in range(4, reader.used_columns[sheet] + 1)
            if reader.effective(sheet, row, column) is not None
        ]
        normalized = {
            "nombre": json_value(name),
            "supervisor": json_value(reader.effective(sheet, row, 2)),
            "auxiliar": json_value(reader.effective(sheet, row, 3)),
            "territorios_solicitados": territories,
        }
        reason = None
        if reader.row_has_formula_error(sheet, row):
            columns = ", ".join(reader.raw_row(sheet, row).get("errores_formula", []))
            normalized = None
            reason = f"error de fórmula en la fila (celdas: {columns})"
        records.append(
            {
                "pestania": sheet,
                "fila": row,
                "rango": f"A{row}:{get_column_letter(reader.used_columns[sheet])}{row}",
                "tipo": "grupo",
                "bruto": reader.raw_row(sheet, row),
                "normalizado": normalized,
                "estado": "conflicto" if reason else "pendiente",
                "motivo": reason,
                "destino_tabla": "grupos_servicio",
            }
        )
    return records


def build_personal_territory_records(reader: WorkbookReader) -> list[dict[str, Any]]:
    sheet = "Grupos - Colores"
    records: list[dict[str, Any]] = []
    for row in range(15, 24):
        name = reader.effective(sheet, row, 1)
        if name is None:
            continue
        normalized = {
            "persona": json_value(name),
            "territorio_bruto": json_value(reader.effective(sheet, row, 2)),
            "activo_bruto": json_value(reader.effective(sheet, row, 4)),
            "periodos": [
                json_value(reader.effective(sheet, row, column))
                for column in range(6, 12)
                if reader.effective(sheet, row, column) is not None
            ],
        }
        reasons: list[str] = []
        active = normalize_text(reader.effective(sheet, row, 4))
        if active == "confirmar":
            reasons.append("estado de asignación personal requiere confirmación")
        if reader.row_has_formula_error(sheet, row):
            columns = ", ".join(reader.raw_row(sheet, row).get("errores_formula", []))
            reasons.append(f"error de fórmula en la fila (celdas: {columns})")
        if reasons:
            normalized = None
        records.append(
            {
                "pestania": sheet,
                "fila": row,
                "rango": f"A{row}:{get_column_letter(reader.used_columns[sheet])}{row}",
                "tipo": "territorio_personal",
                "bruto": reader.raw_row(sheet, row),
                "normalizado": normalized,
                "estado": "conflicto" if reasons else "pendiente",
                "motivo": " ; ".join(reasons) if reasons else None,
                "destino_tabla": "territorio_personal_reservas",
            }
        )
    return records


def build_unattached_comment_records(
    reader: WorkbookReader, existing_records: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """Ningún comentario queda afuera aunque su fila no sea una salida reconocida."""

    covered: set[tuple[str, str]] = set()
    for record in existing_records:
        raw = record.get("bruto") or {}
        raw_sheet = raw.get("hoja")
        raw_row = raw.get("fila")
        for cell in raw.get("celdas") or []:
            if raw_sheet and raw_row and cell.get("comentario"):
                covered.add((raw_sheet, f"{cell['columna']}{raw_row}"))

    records: list[dict[str, Any]] = []
    for worksheet in reader.source.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.comment is None or (worksheet.title, cell.coordinate) in covered:
                    continue
                records.append(
                    {
                        "pestania": worksheet.title,
                        "fila": cell.row,
                        "rango": cell.coordinate,
                        "tipo": "otro",
                        "bruto": reader.raw_with_context(
                            worksheet.title, cell.row, cell=cell.coordinate
                        ),
                        "normalizado": {
                            "comentario_celda": cell.comment.text,
                            "autor_comentario": cell.comment.author,
                            "celda_origen": cell.coordinate,
                        },
                        "estado": "conflicto",
                        "motivo": "comentario fuera de una fila clasificada; requiere revisión",
                        "destino_tabla": None,
                    }
                )
    return records


def build_all_records(reader: WorkbookReader) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    events, agenda_by_sheet = agenda_events(reader)
    records: list[dict[str, Any]] = []
    salida_records = build_salida_records(reader, events)
    records.extend(salida_records)
    # agenda_by_sheet empieza con las filas verticales; las dos filas
    # sintéticas representan los dos bloques horizontales de la misma hoja.
    for sheet, _, _, _ in HORIZONTAL_EXTRA_BLOCKS:
        agenda_by_sheet[sheet] = agenda_by_sheet.get(sheet, 0) + 1
    records.extend(build_history_records(reader))
    records.extend(build_conductor_records(reader, events))
    territory_records, point_records, zoom_count = build_territory_records(reader)
    records.extend(territory_records)
    records.extend(point_records)
    records.extend(build_group_records(reader))
    records.extend(build_personal_territory_records(reader))
    records.extend(build_unattached_comment_records(reader, records))

    # Identidad estable entre corridas: no depende del UUID del lote. El rango
    # exacto distingue subregistros y el hash congela la revisión del archivo.
    for record in records:
        identity = "|".join(
            (
                reader.source_sha256,
                record["pestania"],
                record["rango"] or str(record["fila"]),
                record["tipo"],
            )
        )
        record["source_key"] = hashlib.sha256(identity.encode("utf-8")).hexdigest()
        record["record_role"] = record["tipo"]

    formula_error_cells = 0
    for sheet in reader.source.sheetnames:
        width = reader.used_columns[sheet]
        ws = reader.source[sheet]
        for row in range(1, ws.max_row + 1):
            formula_error_cells += len(reader.formula_error_columns(sheet, row))

    summary: dict[str, Any] = {
        "agenda_por_pestania": dict(sorted(agenda_by_sheet.items())),
        "formula_error_cells": formula_error_cells,
        "zoom_links_excluidos": zoom_count,
        "salidas_sin_booleano": sum(
            1
            for record in records
            if record["tipo"] == "salida"
            and record["estado"] != "conflicto"
            and record["normalizado"] is not None
            and record["normalizado"]["estado"] == "sin_dato"
        ),
        "salidas_false_sin_confirmar": sum(
            1
            for record in records
            if record["tipo"] == "salida"
            and record["normalizado"] is not None
            and record["normalizado"].get("estado_fuente") is False
            and record["normalizado"].get("estado") == "sin_dato"
        ),
        "comentarios_celda": sum(
            1
            for sheet in reader.source.sheetnames
            for row in reader.source[sheet].iter_rows()
            for cell in row
            if cell.comment is not None
        ),
        "historial_textual": sum(
            1
            for record in records
            if record["tipo"] == "historial_territorio"
            and (
                record["normalizado"] is None
                or record["normalizado"].get("clase") in {"bloqueo", "excepcion"}
            )
        ),
    }
    return records, summary


def count_records(records: list[dict[str, Any]]) -> Counter[str]:
    return Counter(record["tipo"] for record in records)


def conflict_counts(records: list[dict[str, Any]]) -> Counter[str]:
    return Counter(
        record["motivo"] or "sin motivo"
        for record in records
        if record["estado"] == "conflicto"
    )


def print_summary(records: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    type_counts = count_records(records)
    tab_type_counts = Counter((record["pestania"], record["tipo"]) for record in records)
    conflicts = conflict_counts(records)

    print("FUENTE")
    print(f"  archivo: {SOURCE_FILE}")
    print(f"  drive_id: {DRIVE_ID}")
    if records:
        print(f"  sha256: {records[0]['bruto']['archivo_sha256']}")
    print(f"  parser_version: {PARSER_VERSION}")
    print(f"  registros staging a cargar: {len(records)}")
    print("TIPOS")
    for kind in sorted(type_counts):
        expected = EXPECTED_COUNTS.get(kind)
        suffix = f" (esperado {expected})" if expected is not None else ""
        print(f"  {kind}: {type_counts[kind]}{suffix}")
    print("POR_PESTANIA_Y_TIPO")
    for (sheet, kind), count in sorted(tab_type_counts.items()):
        print(f"  {sheet!r} | {kind}: {count}")
    print("AGENDA_POR_PESTANIA")
    for sheet, count in sorted(summary["agenda_por_pestania"].items()):
        print(f"  {sheet!r}: {count}")
    print("REGLAS_CONTROLADAS")
    print(f"  salidas sin booleano concluyente: {summary['salidas_sin_booleano']}")
    print(
        "  casillas False preservadas como sin confirmar: "
        f"{summary['salidas_false_sin_confirmar']}"
    )
    print(f"  comentarios de celda preservados: {summary['comentarios_celda']}")
    print(f"  textos del historial no convertidos a fecha: {summary['historial_textual']}")
    print(f"  errores de fórmula visibles en celdas: {summary['formula_error_cells']}")
    print(f"  enlaces Zoom excluidos de punto_encuentro: {summary['zoom_links_excluidos']}")
    print(f"  conflictos: {sum(conflicts.values())}")
    print("MOTIVOS_DE_CONFLICTO")
    for reason, count in conflicts.most_common():
        print(f"  {count} | {reason}")


def validate(records: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    counts = count_records(records)
    problems: list[str] = []
    for kind, expected in EXPECTED_COUNTS.items():
        actual = counts.get(kind, 0)
        if actual != expected:
            problems.append(f"tipo {kind}: actual {actual}, esperado {expected}")

    for sheet, expected in EXPECTED_AGENDA_BY_SHEET.items():
        actual = summary["agenda_por_pestania"].get(sheet, 0)
        if actual != expected:
            problems.append(f"pestaña {sheet!r}: actual {actual}, esperado {expected}")

    if summary["salidas_false_sin_confirmar"] == 0:
        problems.append("no se detectaron casillas False para preservar como sin confirmar")
    if summary["comentarios_celda"] != 11:
        problems.append(
            f"comentarios de celda: actual {summary['comentarios_celda']}, esperado 11"
        )
    if summary["historial_textual"] != 70:
        problems.append(
            f"textos de historial: actual {summary['historial_textual']}, esperado 70"
        )
    if summary["zoom_links_excluidos"] != 1:
        problems.append(
            f"enlaces Zoom excluidos: actual {summary['zoom_links_excluidos']}, esperado 1"
        )

    # Los encabezados quedan fuera de los alias; el valor `-` sí es una
    # variante textual del origen y por eso permanece como alias dudoso.
    aliases = {
        record["normalizado"]["alias"]
        for record in records
        if record["tipo"] == "conductor" and record["normalizado"] is not None
    }
    if len(aliases) != 55:
        problems.append(f"alias de conductor únicos: actual {len(aliases)}, esperado 55")
    if any(
        record["tipo"] == "conductor"
        and (
            record["normalizado"].get("conductor_id") is not None
            or record["normalizado"].get("confianza") != "dudoso"
        )
        for record in records
    ):
        problems.append("hay un alias de conductor fusionado o con confianza distinta de dudoso")

    if problems:
        raise ETLError("Validación detenida; no se escribió staging:\n- " + "\n- ".join(problems))


class SupabaseRest:
    def __init__(self, project_ref: str, key_file: Path):
        if project_ref != DEV_REF or project_ref == PROD_REF:
            raise ETLError(
                f"Ref inseguro {project_ref!r}. Este ETL acepta únicamente el clon {DEV_REF}."
            )
        if not key_file.is_file():
            raise ETLError(f"No existe SUPABASE_SERVICE_ROLE_KEY_FILE: {key_file}")
        self.project_ref = project_ref
        self.base = f"https://{project_ref}.supabase.co/rest/v1"
        self.key = key_file.read_text(encoding="utf-8").strip()
        if not self.key:
            raise ETLError("El archivo de service role key está vacío.")
        self.headers = {
            "apikey": self.key,
            "Authorization": f"Bearer {self.key}",
            "Content-Type": "application/json",
            "Accept": "application/json",
            "Accept-Profile": "public",
        }

    def request(
        self,
        method: str,
        table: str,
        *,
        query: dict[str, str] | None = None,
        body: Any = None,
        prefer: str | None = None,
    ) -> Any:
        url = f"{self.base}/{quote(table, safe='')}"
        if query:
            url += "?" + urlencode(query)
        headers = dict(self.headers)
        if prefer:
            headers["Prefer"] = prefer
        payload = None
        if body is not None:
            payload = json.dumps(body, ensure_ascii=False, allow_nan=False).encode("utf-8")
        request = Request(url, data=payload, headers=headers, method=method)
        try:
            with urlopen(request, timeout=60) as response:
                data = response.read()
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")[:1000]
            raise ETLError(f"Supabase {exc.code} {exc.reason}: {detail}") from exc
        except URLError as exc:
            raise ETLError(f"No se pudo conectar con Supabase: {exc.reason}") from exc
        if not data:
            return None
        try:
            return json.loads(data.decode("utf-8"))
        except json.JSONDecodeError:
            return data.decode("utf-8", errors="replace")

    def existing_importations(self, source_sha256: str) -> list[dict[str, Any]]:
        return self.request(
            "GET",
            "importaciones",
            query={
                "source_sha256": f"eq.{source_sha256}",
                "parser_version": f"eq.{PARSER_VERSION}",
                "select": "id,estado,corrida_at,source_sha256,parser_version",
            },
        ) or []

    def create_importation(self, reader: WorkbookReader) -> str:
        result = self.request(
            "POST",
            "importaciones",
            body={
                "archivo": SOURCE_FILE,
                "drive_id": DRIVE_ID,
                "source_sha256": reader.source_sha256,
                "source_size_bytes": reader.source_size_bytes,
                "parser_version": PARSER_VERSION,
                "nota": (
                    "ETL de staging solamente. Fuente leída en modo solo lectura desde Drive. "
                    "No aplica salidas, resultados, historial ni conductores; requiere revisión."
                ),
                "estado": "en_curso",
            },
            prefer="return=representation",
        )
        if not isinstance(result, list) or not result or not result[0].get("id"):
            raise ETLError("Supabase no devolvió el id de importaciones.")
        return str(result[0]["id"])

    def insert_records(self, importation_id: str, records: list[dict[str, Any]]) -> int:
        payload_records = []
        for record in records:
            payload_records.append(
                {
                    "importacion_id": importation_id,
                    "pestania": record["pestania"],
                    "fila": record["fila"],
                    "rango": record["rango"],
                    "tipo": record["tipo"],
                    "bruto": record["bruto"],
                    "normalizado": record["normalizado"],
                    "estado": record["estado"],
                    "motivo": record["motivo"],
                    "destino_tabla": record["destino_tabla"],
                    "source_key": record["source_key"],
                    "record_role": record["record_role"],
                    "resolution_status": "open",
                    "quality_status": (
                        "blocked" if record["estado"] == "conflicto" else "warning"
                    ),
                }
            )

        inserted = 0
        for start in range(0, len(payload_records), 100):
            batch = payload_records[start : start + 100]
            self.request(
                "POST",
                "importacion_registros",
                body=batch,
                prefer="return=minimal",
            )
            inserted += len(batch)
            print(f"  staging: {inserted}/{len(payload_records)} registros")
        return inserted

    def mark_reverted(self, importation_id: str) -> None:
        self.request(
            "PATCH",
            "importaciones",
            query={"id": f"eq.{importation_id}"},
            body={"estado": "revertida"},
            prefer="return=minimal",
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        default="/tmp/Salidas-de-predicacion.xlsx",
        help="ruta local de la copia XLSX leída desde Drive",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="clasifica y valida sin conectarse ni escribir Supabase",
    )
    parser.add_argument(
        "--allow-existing",
        action="store_true",
        help=(
            "permite reintentar el mismo hash/parser únicamente si todas las "
            "corridas anteriores están revertidas"
        ),
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = Path(args.input).expanduser().resolve()
    reader = WorkbookReader(input_path)
    try:
        records, summary = build_all_records(reader)
        print_summary(records, summary)
        validate(records, summary)
        print("VALIDACION: OK; los totales y las reglas de seguridad coinciden.")

        if args.dry_run:
            print("DRY-RUN: no se escribió ninguna tabla.")
            return 0

        project_ref = os.environ.get("SUPABASE_PROJECT_REF")
        key_file_value = os.environ.get("SUPABASE_SERVICE_ROLE_KEY_FILE")
        if not project_ref or not key_file_value:
            raise ETLError(
                "Faltan SUPABASE_PROJECT_REF y SUPABASE_SERVICE_ROLE_KEY_FILE."
            )
        # La única ref admitida es el clon exacto. Esto evita que un typo o un
        # ref viejo redirija el ETL a otro proyecto.
        if project_ref != DEV_REF or project_ref == PROD_REF:
            raise ETLError(
                f"Ref inseguro {project_ref!r}; se esperaba únicamente el clon {DEV_REF}."
            )
        api = SupabaseRest(project_ref, Path(key_file_value).expanduser().resolve())
        existing = api.existing_importations(reader.source_sha256)
        active_existing = [
            item for item in existing if item.get("estado") != "revertida"
        ]
        if active_existing:
            states = ", ".join(
                f"{item.get('id')} ({item.get('estado')})"
                for item in active_existing
            )
            raise ETLError(
                "Ya existe una corrida activa para este hash y versión del parser: "
                f"{states}. Debe resolverse o revertirse antes de reintentar."
            )
        if existing and not args.allow_existing:
            states = ", ".join(
                f"{item.get('id')} ({item.get('estado')})" for item in existing
            )
            raise ETLError(
                "Ya existe una corrida revertida para este hash y versión del parser: "
                f"{states}. Usar --allow-existing sólo para crear un reintento auditable."
            )

        importation_id = api.create_importation(reader)
        print(f"importaciones: corrida creada en el clon ({importation_id})")
        try:
            inserted = api.insert_records(importation_id, records)
        except Exception:
            try:
                api.mark_reverted(importation_id)
                print("importaciones: corrida parcial marcada como revertida")
            except Exception as revert_error:
                print(
                    f"ADVERTENCIA: no se pudo marcar la corrida como revertida: {revert_error}",
                    file=sys.stderr,
                )
            raise
        print(f"CARGA: {inserted} registros escritos únicamente en staging.")
        print(f"CARGA: importaciones.id={importation_id}, estado=en_curso")
        return 0
    finally:
        reader.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except ETLError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
